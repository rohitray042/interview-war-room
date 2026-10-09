import copy
from io import BytesIO

import pytest
from alembic import command
from docx import Document as WordDocument
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.contracts import LLMResult
from app.db import migration_config
from app.document_parser import MAX_BYTES, DocumentError, extract_file
from app.main import create_app

RESUME = """Test Candidate
SUMMARY
Data Engineer with 4 years of experience.
SKILLS
Python, Snowflake, SQL
EXPERIENCE
Engineer | Example Company | 2022 - Present
- Built Snowflake streams and tasks for ingestion.
- Reduced SQL query time by 35%.
EDUCATION
BTech Computer Science
CERTIFICATIONS
SnowPro Core
"""
JD = """Senior Data Engineer
Must have
Snowflake
Kafka
Python
Good to have
AWS
Nice to have
Databricks
Responsibilities
You will design data pipelines and collaborate with stakeholders.
Experience: 4+ years
"""


@pytest.fixture
def client(settings):
    with TestClient(create_app(settings)) as value:
        yield value


def paste(client, kind, text):
    response = client.post("/api/v1/documents/text", json={"kind": kind, "text": text})
    assert response.status_code == 200, response.text
    return response.json()


def confirm(client, bundle):
    analysis = bundle["analyses"][0]
    response = client.post(
        f"/api/v1/analyses/{analysis['id']}/confirm", json={"revision": analysis["revision"]}
    )
    assert response.status_code == 200, response.text
    return response.json()


def pdf_bytes(text="Test Candidate"):
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
    )
    stream = DecodedStreamObject()
    stream.set_data(f"BT /F1 12 Tf 50 700 Td ({text}) Tj ET".encode())
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def test_pdf_upload(client):
    response = client.post(
        "/api/v1/documents/upload",
        data={"kind": "resume"},
        files={
            "file": ("resume.pdf", pdf_bytes("Test Candidate SQL Snowflake"), "application/pdf")
        },
    )
    assert response.status_code == 200, response.text
    assert "Snowflake" in response.json()["document"]["text"]
    assert response.json()["document"]["warnings"]


def test_docx_upload_preserves_paragraph_table_order(client):
    doc = WordDocument()
    doc.add_paragraph("Test Candidate")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Skills"
    table.cell(0, 1).text = "Python"
    doc.add_paragraph("Education")
    output = BytesIO()
    doc.save(output)
    response = client.post(
        "/api/v1/documents/upload",
        data={"kind": "jd"},
        files={"file": ("role.docx", output.getvalue())},
    )
    assert response.status_code == 200
    assert response.json()["document"]["text"].index("Python") < response.json()["document"][
        "text"
    ].index("Education")


@pytest.mark.parametrize(
    "filename,data,status",
    [
        ("empty.txt", b"", 422),
        ("empty.txt", b"  \n", 422),
        ("image.png", b"image", 415),
        ("broken.pdf", b"%PDF-invalid", 422),
        ("broken.docx", b"PK-invalid", 422),
        ("binary.txt", b"\xff\xfe\xfd", 422),
        ("large.txt", b"x" * (MAX_BYTES + 1), 413),
        ("text.txt", b"x" * 100001, 413),
    ],
)
def test_bad_uploads(client, filename, data, status):
    result = client.post(
        "/api/v1/documents/upload", data={"kind": "resume"}, files={"file": (filename, data)}
    )
    assert result.status_code == status, result.text
    assert client.get("/api/v1/documents").json() == []


def test_blank_pdf_rejected():
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    output = BytesIO()
    writer.write(output)
    with pytest.raises(DocumentError, match="No readable text"):
        extract_file("scan.pdf", output.getvalue())


def test_missing_sections_and_explicit_tiers(client):
    bundle = paste(client, "resume", "Candidate Name\nSQL")
    assert bundle["analyses"][0]["warnings"]
    assert not any(c["category"] == "years_experience" for c in bundle["analyses"][0]["claims"])
    jd = paste(client, "jd", "Python and Kafka")
    assert all(c["requirement"] == "unspecified" for c in jd["analyses"][0]["claims"])
    jd2 = paste(client, "jd", JD)
    tiers = {c["value"]: c["requirement"] for c in jd2["analyses"][0]["claims"]}
    assert tiers["Kafka"] == "must_have"
    assert tiers["AWS"] == "good_to_have"
    assert tiers["Databricks"] == "nice_to_have"


def test_mixed_tiers_are_flagged_not_guessed(client):
    bundle = paste(client, "jd", "Python required; Kafka preferred.")
    assert all(c["uncertain"] for c in bundle["analyses"][0]["claims"])
    assert all(c["requirement"] == "unspecified" for c in bundle["analyses"][0]["claims"])


def test_ai_cannot_claim_java_from_javascript(client):
    bundle = paste(client, "resume", "Candidate Name\nSKILLS\nJavaScript")
    source = bundle["document"]
    start = source["text"].index("JavaScript")
    client.app.state.llm = FakeExtractor(
        {
            "claims": [
                {
                    "category": "programming_languages",
                    "value": "Java",
                    "evidence": [
                        {
                            "document_id": source["id"],
                            "start": start,
                            "end": start + 10,
                            "quote": "JavaScript",
                        }
                    ],
                }
            ]
        }
    )
    assert client.post(f"/api/v1/documents/{source['id']}/ai-analysis").status_code == 422


def test_evidence_offsets_duplicate_uploads_and_kind_separation(client):
    first = paste(client, "resume", RESUME)
    for claim in first["analyses"][0]["claims"]:
        assert claim["evidence"]
        for e in claim["evidence"]:
            assert e["document_id"] == first["document"]["id"]
            assert first["document"]["text"][e["start"] : e["end"]] == e["quote"]
    confirm(client, first)
    duplicate = client.post(
        "/api/v1/documents/upload",
        data={"kind": "resume"},
        files={"file": ("renamed.txt", RESUME.encode())},
    ).json()
    assert duplicate["duplicate"] is True
    assert duplicate["document"]["id"] == first["document"]["id"]
    assert duplicate["analyses"][0]["status"] == "confirmed"
    assert paste(client, "jd", RESUME)["document"]["id"] != first["document"]["id"]


def test_corrections_delete_add_confirm_and_conflict(client):
    bundle = paste(client, "resume", RESUME)
    analysis = bundle["analyses"][0]
    claims = analysis["claims"][:-1]
    claims[0]["value"] = "Corrected Candidate"
    claims.append(
        {
            "id": "manual",
            "category": "skills",
            "value": "Rust",
            "origin": "extracted",
            "reviewed": False,
            "uncertain": False,
            "requirement": "unspecified",
            "evidence": [],
        }
    )
    url = f"/api/v1/analyses/{analysis['id']}"
    response = client.put(url, json={"revision": 1, "claims": claims})
    assert response.status_code == 200, response.text
    saved = response.json()
    assert saved["claims"][0]["origin"] == "user"
    assert saved["claims"][0]["evidence"] == []
    assert saved["claims"][-1]["origin"] == "user"
    assert client.put(url, json={"revision": 1, "claims": claims}).status_code == 409
    assert client.post(url + "/confirm", json={"revision": 1}).status_code == 409
    assert client.post(url + "/confirm", json={"revision": 2}).json()["status"] == "confirmed"
    assert client.put(url, json={"revision": 3, "claims": claims}).status_code == 409
    draft = client.post(url + "/revision").json()
    assert draft["id"] != analysis["id"] and draft["status"] == "draft"


def test_invalid_evidence_is_rejected(client):
    a = paste(client, "resume", RESUME)["analyses"][0]
    a["claims"][0]["evidence"][0]["quote"] = "Fabricated achievement"
    assert (
        client.put(
            f"/api/v1/analyses/{a['id']}", json={"revision": 1, "claims": a["claims"]}
        ).status_code
        == 422
    )


def test_comparison_gated_and_evidence_based(client):
    r = paste(client, "resume", RESUME)
    j = paste(client, "jd", JD)
    payload = {"resume_id": r["analyses"][0]["id"], "jd_id": j["analyses"][0]["id"]}
    assert client.post("/api/v1/targets", json=payload).status_code == 409
    confirm(client, r)
    confirm(client, j)
    result = client.post("/api/v1/targets", json=payload)
    assert result.status_code == 200, result.text
    items = {i["topic"]: i for i in result.json()["result"]["items"]}
    assert items["Snowflake"]["status"] == "strong"
    assert items["Kafka"]["status"] == "missing"
    assert items["Python"]["status"] == "partial"
    assert items["Kafka"]["priority"] == "high"
    assert items["Databricks"]["priority"] == "low"
    assert items["Kafka"]["jd_claims"][0]["evidence"]
    assert client.post("/api/v1/targets", json=payload).json()["id"] == result.json()["id"]


@pytest.mark.parametrize("kind", ["resume", "jd"])
def test_unsupported_user_correction_blocks_comparison(client, kind):
    bundles = {"resume": paste(client, "resume", RESUME), "jd": paste(client, "jd", JD)}
    analysis = bundles[kind]["analyses"][0]
    claim = copy.deepcopy(analysis["claims"][0])
    claim.update(id="unsupported", category="skills", value="Rust", evidence=[])
    response = client.put(
        f"/api/v1/analyses/{analysis['id']}",
        json={"revision": analysis["revision"], "claims": analysis["claims"] + [claim]},
    )
    assert response.status_code == 200
    bundles[kind]["analyses"][0] = response.json()
    r, j = confirm(client, bundles["resume"]), confirm(client, bundles["jd"])
    response = client.post("/api/v1/targets", json={"resume_id": r["id"], "jd_id": j["id"]})
    assert response.status_code == 422
    assert "exact source evidence" in response.json()["detail"]


def test_deleted_mentions_cannot_be_classified_missing(client):
    bundle = paste(client, "resume", RESUME)
    analysis = bundle["analyses"][0]
    response = client.put(
        f"/api/v1/analyses/{analysis['id']}",
        json={"revision": 1, "claims": [c for c in analysis["claims"] if c["value"] != "Python"]},
    )
    bundle["analyses"][0] = response.json()
    r = confirm(client, bundle)
    j = confirm(client, paste(client, "jd", JD))
    result = client.post("/api/v1/targets", json={"resume_id": r["id"], "jd_id": j["id"]}).json()
    rows = {row["topic"]: row for row in result["result"]["items"]}
    assert rows["Python"]["status"] == "needs_revision"
    assert rows["Python"]["source_mentions"][0]["quote"] == "Python, Snowflake, SQL"
    assert rows["Kafka"]["absence_check"]
    assert rows["Kafka"]["source_mentions"] == []
    for row in rows.values():
        for side, source in (("resume_claims", RESUME), ("jd_claims", JD)):
            for claim in row[side]:
                assert claim["evidence"]
                for evidence in claim["evidence"]:
                    assert source[evidence["start"] : evidence["end"]] == evidence["quote"]


def test_negated_experience_not_a_strong_match(client):
    r = confirm(
        client,
        paste(
            client,
            "resume",
            "Candidate Name\nEXPERIENCE\n- No experience with Kafka; learning Python.",
        ),
    )
    j = confirm(client, paste(client, "jd", "Must have: Kafka"))
    target = client.post("/api/v1/targets", json={"resume_id": r["id"], "jd_id": j["id"]}).json()
    assert target["result"]["items"][0]["status"] == "needs_revision"


class FakeExtractor:
    def __init__(self, data):
        self.data = data

    async def generate_structured(self, request):
        return LLMResult(data=self.data, provider="fake", model="fixture", request_id="fake-1")


def test_llm_unavailable_keeps_local_analysis(client):
    bundle = paste(client, "resume", RESUME)
    assert (
        client.post(f"/api/v1/documents/{bundle['document']['id']}/ai-analysis").status_code == 503
    )
    assert len(client.get("/api/v1/documents/" + bundle["document"]["id"]).json()["analyses"]) == 1


@pytest.mark.parametrize(
    "mode", ["invalid_schema", "invented_value", "invented_quote", "wrong_document"]
)
def test_bad_llm_output_rejected(client, mode):
    bundle = paste(client, "resume", RESUME)
    claim = copy.deepcopy(
        next(c for c in bundle["analyses"][0]["claims"] if c["value"] == "Python")
    )
    if mode == "invented_value":
        claim["value"] = "Invented billion dollar achievement"
    if mode == "invented_quote":
        claim["evidence"][0]["quote"] = "Invented source"
    if mode == "wrong_document":
        claim["evidence"][0]["document_id"] = "other-document"
    data = {"claims": [claim]} if mode != "invalid_schema" else {"claims": [], "made_up_score": 99}
    client.app.state.llm = FakeExtractor(data)
    assert (
        client.post(f"/api/v1/documents/{bundle['document']['id']}/ai-analysis").status_code == 422
    )


def test_ai_claims_need_individual_review(client):
    bundle = paste(client, "resume", RESUME)
    claim = next(c for c in bundle["analyses"][0]["claims"] if c["value"] == "Python")
    client.app.state.llm = FakeExtractor({"claims": [claim]})
    result = client.post(f"/api/v1/documents/{bundle['document']['id']}/ai-analysis").json()
    assert result["claims"][0]["origin"] == "ai_inferred"
    url = f"/api/v1/analyses/{result['id']}"
    assert client.post(url + "/confirm", json={"revision": 1}).status_code == 422
    result["claims"][0]["reviewed"] = True
    assert client.put(url, json={"revision": 1, "claims": result["claims"]}).status_code == 200
    assert client.post(url + "/confirm", json={"revision": 2}).status_code == 200


def test_analyzer_persistence(settings):
    with TestClient(create_app(settings)) as client:
        r = confirm(client, paste(client, "resume", RESUME))
        j = confirm(client, paste(client, "jd", JD))
        target = client.post(
            "/api/v1/targets", json={"resume_id": r["id"], "jd_id": j["id"]}
        ).json()
    with TestClient(create_app(settings)) as restarted:
        assert len(restarted.get("/api/v1/documents").json()) == 2
        assert restarted.get("/api/v1/targets").json()[0]["id"] == target["id"]
        data = restarted.get("/api/v1/documents/" + r["document_id"]).json()
        assert data["analyses"][0]["status"] == "confirmed"
        assert data["analyses"][0]["claims"][0]["evidence"]


def test_migration_preserves_profile(settings):
    config = migration_config(settings)
    command.downgrade(config, "0001")
    import sqlite3

    with sqlite3.connect(settings.resolved_database_path) as db:
        db.execute(
            "INSERT INTO profiles VALUES ('profile',1,'Preserved','','Data Engineer',"
            "'[]','[]',60,'2026-09-26','2026-09-26')"
        )
    command.upgrade(config, "head")
    with TestClient(create_app(settings)) as client:
        assert client.get("/api/v1/profile").json()["name"] == "Preserved"
